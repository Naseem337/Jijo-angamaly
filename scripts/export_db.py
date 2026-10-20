"""Export the bot's local database to an Excel file you can open normally.

    .venv\\Scripts\\python.exe scripts\\export_db.py            → data\\database_export.xlsx

One sheet per table (Messages first), times shown in India time, JSON columns readable.
The file holds customers' phone numbers and chats: keep it private (data\\ is not in git).
Needs openpyxl (pip install openpyxl).
"""

from __future__ import annotations

import json
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.core.clock import to_ist  # noqa: E402
from app.settings import get_settings  # noqa: E402

# Most useful first; the rest follow alphabetically.
ORDER = ["messages", "leads", "enquiries", "customers", "requirements", "alerts", "counters"]
HIDE = {"processed_events", "cached_media", "seen_cars", "pending_actions"}  # bot bookkeeping
TIME_COLUMNS = {"ts", "created_at", "updated_at", "expires_at", "consent_at", "opted_out_at", "paused_at",
                "last_inbound_at", "last_outbound_at", "still_looking_sent_at", "last_near_match_at"}


def readable(column: str, value):
    if value is None:
        return ""
    if column in TIME_COLUMNS and isinstance(value, str):
        try:
            return to_ist(datetime.fromisoformat(value)).strftime("%d %b %Y %H:%M")
        except ValueError:
            return value
    if column in ("meta", "requirement", "args") and isinstance(value, str):
        try:
            data = json.loads(value)
        except ValueError:
            return value
        return ", ".join(f"{k}: {v}" for k, v in data.items() if v not in (None, "", [], {})) if isinstance(data, dict) else value
    return value


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    db = Path(get_settings().sqlite_path)
    if not db.is_absolute():
        db = ROOT / db
    out = db.parent / "database_export.xlsx"
    con = sqlite3.connect(db)
    tables = [r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name")]
    tables = [t for t in ORDER if t in tables] + [t for t in tables if t not in ORDER and t not in HIDE]

    book = Workbook()
    book.remove(book.active)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for table in tables:
        cursor = con.execute(f"SELECT * FROM {table} ORDER BY rowid")
        columns = [d[0] for d in cursor.description]
        rows = cursor.fetchall()
        sheet = book.create_sheet(table.replace("_", " ").title()[:31])
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True), header_fill
        for row in rows:
            sheet.append([readable(c, v) for c, v in zip(columns, row)])
        for i, column in enumerate(columns, 1):
            width = max([len(str(column))] + [len(str(sheet.cell(r, i).value or "")) for r in range(2, min(len(rows), 200) + 2)])
            sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        print(f"{table:<17}{len(rows):>5} rows")
    book.save(out)
    print(f"\nSaved: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
